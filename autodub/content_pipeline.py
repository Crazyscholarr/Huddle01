"""Pipeline biến kho truyện thô thành kế hoạch sản xuất YouTube.

Đầu vào tương thích trực tiếp JSON do ``pixel-acre-brook-comet`` xuất ra và
nhận cả SQLite/JSON có tên cột phổ biến.  Phân tích luôn có lớp heuristic chạy
offline; khi provider API trong ``config.yaml`` sẵn sàng, một lượt AI sẽ bổ
sung tiêu đề, mô tả, thumbnail, outline và chấm điểm sâu hơn.
"""
from __future__ import annotations

import asyncio
import hashlib
import html
import importlib.util
import json
import os
import re
import sqlite3
import sys
import tempfile
import threading
import time
import zipfile
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_DB = os.path.join(ROOT, "data", "content_ideas.sqlite")
DEFAULT_OUTPUT_DIR = os.path.join(ROOT, "output", "content_plans")

TITLE_PROMPT = """Bạn đặt tiêu đề video cho kênh audio Việt Nam \"Gốc Mít Kể Chuyện\",
chủ đề chuyện gia đình, tuổi già và chuyện tâm linh ông bà kể, khán giả từ 45 tuổi trở lên.

NỘI DUNG TRUYỆN: [DÁN TÓM TẮT]

Hãy tạo đúng 8 tiêu đề theo công thức 3 phần:
[MÓC CẢM XÚC] : [MỆNH ĐỀ VIẾT HOA TOÀN BỘ] | [ĐUÔI TỪ KHÓA]

Móc cảm xúc chỉ chọn trong: Nghe Mà Thấm / Nghe THẤM Tận Xương / Nghe Là Khóc /
Nghe Mà Nghẹn Lòng / Nghe Mà Rơi Nước Mắt / Nghe Sướng Lỗ Tai / Nghe Mà Sốc /
Truyện Ngắn Tuổi Xế Chiều Cực Hay. Riêng chuyện tâm linh có thể chọn: Nghe Mà
Lạnh Gáy / Chuyện Ông Bà Kể / Nghe Mà Rùng Mình / Chuyện Làng Quê Kỳ Bí.

Mệnh đề viết hoa phải nêu đúng tình huống sốc nhất, dài 8 đến 14 từ, viết hoa
toàn bộ, là một sự việc cụ thể chứ không phải cảm xúc chung. Có con số càng tốt.

Đuôi từ khóa chỉ chọn trong: Kể Chuyện Đêm Khuya / Đọc Truyện Đêm Khuya /
Kể Chuyện Tuổi Già / Kể Chuyện Làng Quê / Kể Chuyện Tâm Linh.

Tổng mỗi tiêu đề không quá 100 ký tự; không đưa tên kênh hoặc số tập vào tiêu đề;
không tiết lộ kết thúc; không dịch hoặc sao chép tiêu đề nguồn."""

DESCRIPTION_PROMPT = """Tạo 3 mô tả YouTube tiếng Việt cho truyện audio gia đình,
mỗi mô tả 70-120 từ. Mở bằng tình huống gợi tò mò, nêu giá trị cảm xúc, mời
khán giả bình luận/đăng ký tự nhiên và kết thúc bằng 3-5 hashtag phù hợp. Không
bịa đây là chuyện thật và không tiết lộ kết thúc."""

ANALYSIS_PROMPT = """Bạn là biên tập viên chiến lược nội dung YouTube Việt Nam,
chuyên audio chuyện gia đình, tuổi già và chuyện tâm linh ông bà kể cho khán giả
45+. Chỉ lấy mô-típ và
tình huống để sáng tạo lại; không dịch hay sao chép nguyên văn nguồn. Với nguồn
Trung Quốc, hãy đọc để RÚT CHẤT LIỆU KỊCH TÍNH, sau đó diễn đạt toàn bộ kết quả
bằng tiếng Việt. Tuyệt đối không kể lại tuần tự hoặc đưa nguyên văn truyện nguồn
vào hồ sơ sáng tác. Với chuyện tâm linh, thể hiện như ký ức/truyền miệng của ông
bà, không khẳng định là sự thật khoa học, không cổ súy mê tín hoặc hành động nguy
hiểm; trọng tâm vẫn là tình thân, nhân quả và nếp sống làng quê Việt.

Hãy trả về MỘT JSON thuần, không markdown, theo đúng schema:
{
  "title_localized":"Tên làm việc tiếng Việt, cụ thể, không phải bản dịch từng chữ",
  "primary_genre":"...", "emotion":"...", "hook_score":1,
  "plot_twist_score":1, "themes":["..."], "archetypes":["..."],
  "main_hook":"Một câu mô tả móc mở đầu đáng giữ lại",
  "high_tension_scenes":["5-7 cảnh/xung đột kịch tính đáng khai thác"],
  "plot_twists":["2-4 cú lật hoặc bí mật đáng khai thác"],
  "must_change":["tên", "địa danh", "quan hệ", "số liệu", "diễn biến", "lời văn"],
  "rewrite_brief":"Hồ sơ sáng tác tiếng Việt 500-900 từ: tiền đề mới, xung đột, các nhịp leo thang, hook, cú lật và hướng nhân quả; chỉ giữ chất liệu, yêu cầu viết truyện Việt hoàn toàn mới",
  "keywords":["..."], "tags":["..."], "series":"...",
  "titles":[đúng 8 chuỗi], "descriptions":[3 chuỗi],
  "thumbnails":[{"description":"...","colors":"...","text":"..."}, ...3],
  "outline":"Phần 1: ...\nPhần 2: ...\nPhần 3: ...",
  "main_characters":["..."], "production_ease":"Cao|Trung|Thấp",
  "recommendation":"Nên làm|Chờ đợi|Không nên làm",
  "best_publish_time":"..."
}

Điểm Hook/Plot Twist là số nguyên 1-10. Tiêu đề tuân thủ prompt sau; trong lượt
này hãy thay [DÁN TÓM TẮT] bằng chính rewrite_brief vừa rút ra:
""" + TITLE_PROMPT + "\n\nMô tả tuân thủ:\n" + DESCRIPTION_PROMPT


FIELD_ALIASES: Dict[str, Tuple[str, ...]] = {
    "id": ("id", "story_id", "uuid"),
    "title_original": (
        "titleOriginal", "title_original", "original_title", "title", "name",
        "tieu_de_goc", "tiêu đề gốc"),
    "title_localized": (
        "titleLocalized", "title_localized", "localized_title", "title_vi",
        "tieu_de_viet", "tiêu đề việt"),
    "content": (
        "localizedContent", "localized_content", "cleanedContent",
        "cleaned_content", "rawContent", "raw_content", "content", "body",
        "text", "story", "noi_dung", "nội dung"),
    "raw_content": (
        "rawContent", "raw_content", "content", "body", "text", "story",
        "noi_dung", "nội dung"),
    "source": ("source", "source_name", "platform", "nguon", "nguồn"),
    "source_url": ("sourceUrl", "source_url", "url", "link", "permalink"),
    "author": ("author", "username", "user", "tac_gia", "tác giả"),
    "published_at": (
        "publishedAt", "published_at", "createdAt", "created_at", "date",
        "ngay_dang", "ngày đăng"),
    "language": ("language", "lang", "ngon_ngu", "ngôn ngữ"),
    "tags": ("tags", "tag", "labels", "keywords"),
    "notes": ("notes", "note", "crawlNote", "crawl_note", "ghi_chu"),
}


THEME_TERMS: Dict[str, Tuple[str, ...]] = {
    "Tâm linh ông bà kể": (
        "tâm linh", "chuyện ma", "ông bà kể", "người âm", "báo mộng",
        "điềm báo", "miếu", "nghĩa địa", "灵异", "鬼故事", "托梦",
        "民间怪谈", "乡村怪谈"),
    "Mẹ chồng nàng dâu": (
        "mẹ chồng", "nàng dâu", "婆媳", "mother-in-law", "daughter-in-law"),
    "Con cái bất hiếu": (
        "bất hiếu", "không nuôi", "đuổi bố", "đuổi mẹ", "赡养", "不孝",
        "abandoned parent"),
    "Tình yêu xế chiều": (
        "tái hôn", "tuổi già", "xế chiều", "老年", "中老年", "再婚",
        "widow", "widower"),
    "Tranh chấp thừa kế": (
        "thừa kế", "di chúc", "chia đất", "gia sản", "遗产", "分家",
        "inheritance"),
    "Nghiệp báo ngoại tình": (
        "ngoại tình", "bồ nhí", "phản bội", "出轨", "affair", "cheated"),
    "Cô đơn tuổi già": (
        "cô đơn", "viện dưỡng lão", "sống một mình", "空巢", "养老院",
        "nursing home"),
    "Trọng nam khinh nữ": (
        "trọng nam", "khinh nữ", "重男轻女", "son preference"),
    "Giúp việc và chủ nhà": (
        "giúp việc", "người ở", "bảo mẫu", "保姆", "housekeeper"),
    "Mâu thuẫn con rể mẹ vợ": (
        "con rể", "mẹ vợ", "女婿", "岳母", "son-in-law"),
    "Gia đình và lòng tham": (
        "tham lam", "tiền bạc", "chiếm nhà", "lừa tiền", "贪", "money",
        "greed"),
}

EMOTION_TERMS: Dict[str, Tuple[str, ...]] = {
    "Rùng mình": ("tâm linh", "chuyện ma", "báo mộng", "điềm báo", "灵异", "鬼"),
    "Phẫn nộ": ("bất hiếu", "phản bội", "đuổi", "lừa", "đánh", "cướp", "出轨"),
    "Xót xa": ("cô đơn", "khóc", "mất", "bệnh", "nghèo", "qua đời", "孤独"),
    "Nghẹn lòng": ("hy sinh", "nhịn", "chịu đựng", "ân hận", "hối hận"),
    "Hạnh phúc": ("đoàn tụ", "tha thứ", "bình yên", "hạnh phúc", "sum họp"),
    "Sốc": ("bí mật", "sự thật", "không ngờ", "bất ngờ", "真相", "秘密"),
}

ARCHETYPE_TERMS: Dict[str, Tuple[str, ...]] = {
    "Người ông kể chuyện xưa": ("ông kể", "ông nội", "ông ngoại", "老爷爷", "老人讲"),
    "Người bà giữ bí mật làng": ("bà kể", "bà nội", "bà ngoại", "奶奶", "婆婆讲"),
    "Người mẹ chồng khắc nghiệt": ("mẹ chồng", "婆婆"),
    "Nàng dâu chịu thương chịu khó": ("nàng dâu", "con dâu", "儿媳"),
    "Đứa con bất hiếu": ("bất hiếu", "không nuôi", "不孝"),
    "Người mẹ già hy sinh": ("mẹ già", "bà lão", "母亲", "老人"),
    "Người chồng phản bội": ("chồng ngoại tình", "bồ nhí", "丈夫出轨"),
    "Người vợ nhẫn nhịn": ("vợ", "chịu đựng", "妻子"),
    "Kẻ tham gia sản": ("thừa kế", "di chúc", "gia sản", "遗产"),
    "Người hàng xóm biết chuyện": ("hàng xóm", "邻居"),
}

TWIST_TERMS = (
    "bí mật", "sự thật", "hóa ra", "không ngờ", "bất ngờ", "di chúc",
    "xét nghiệm", "đứa con", "đổi tráo", "lật mặt", "真相", "原来", "秘密",
    "báo mộng", "điềm báo", "người âm", "托梦", "灵异",
    "turns out", "secret", "revealed")
HOOK_TERMS = (
    "đuổi", "ngoại tình", "thừa kế", "mất tích", "ly hôn", "đám tang",
    "tái hôn", "bất hiếu", "tranh chấp", "cướp", "lừa", "报应", "出轨",
    "chuyện ma", "tâm linh", "báo mộng", "nghĩa địa", "灵异", "鬼故事",
    "AITA", "inheritance", "divorce")


def clean_text(value: Any) -> str:
    text = html.unescape(str(value or ""))
    text = re.sub(r"<(script|style|noscript|iframe)[^>]*>[\s\S]*?</\1>", " ", text,
                  flags=re.I)
    text = re.sub(r"<br\s*/?>|</(?:p|div|li|h[1-6])>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _lookup(row: Dict[str, Any], field: str, default: Any = "") -> Any:
    folded = {str(k).casefold(): v for k, v in row.items()}
    for alias in FIELD_ALIASES[field]:
        if alias in row and row[alias] not in (None, ""):
            return row[alias]
        value = folded.get(alias.casefold())
        if value not in (None, ""):
            return value
    return default


def _list_value(value: Any) -> List[str]:
    if isinstance(value, (list, tuple, set)):
        return [str(x).strip() for x in value if str(x).strip()]
    if isinstance(value, dict):
        return [str(x).strip() for x in value.values() if str(x).strip()]
    return [x.strip() for x in re.split(r"[,;|\n]+", str(value or "")) if x.strip()]


def detect_language(text: str) -> str:
    sample = str(text or "")[:2000]
    cjk = len(re.findall(r"[\u4e00-\u9fff]", sample))
    vi = len(re.findall(
        r"[ăâêôơưđáàảãạấầẩẫậắằẳẵặéèẻẽẹếềểễệíìỉĩịóòỏõọốồỗộớờởỡợúùủũụứừửữựýỳỷỹỵ]",
        sample, flags=re.I))
    latin = len(re.findall(r"[A-Za-z]", sample))
    if cjk > 12 and cjk > vi:
        return "zh"
    if vi > 6:
        return "vi"
    return "en" if latin > 20 else "vi"


def count_words(text: str, language: str = "") -> int:
    lang = language or detect_language(text)
    if lang == "zh":
        # 1 chữ Hán thường nở thành khoảng 1.3-1.6 từ khi Việt hóa.
        cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
        latin = len(re.findall(r"\b\w+\b", text, flags=re.UNICODE))
        return max(latin, int(cjk * 1.45))
    return len(re.findall(r"\b\w+\b", text, flags=re.UNICODE))


def _stable_id(source: str, url: str, title: str, content: str) -> str:
    digest = hashlib.sha1(
        (source + "|" + url + "|" + title + "|" + content[:500]).encode("utf-8", "ignore")
    ).hexdigest()[:16]
    return "idea_" + digest


def canonical_source_url(value: str) -> str:
    """Chuẩn hóa URL nguồn để nhận ra cùng một chuyện qua link có tracking."""
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        parts = urlsplit(raw)
        if not parts.hostname:
            return raw.rstrip("/")
        blocked = {"spm_id_from", "vd_source", "from", "source", "share_code"}
        query = [(key, val) for key, val in parse_qsl(parts.query, keep_blank_values=False)
                 if not key.lower().startswith("utm_") and key.lower() not in blocked]
        host = (parts.hostname or "").lower()
        if parts.port:
            host += ":" + str(parts.port)
        path = re.sub(r"/{2,}", "/", parts.path or "/").rstrip("/") or "/"
        return urlunsplit((parts.scheme.lower() or "https", host, path,
                           urlencode(query, doseq=True), ""))
    except Exception:
        return raw.split("#", 1)[0].rstrip("/")


def normalize_record(row: Dict[str, Any], index: int = 0) -> Dict[str, Any]:
    title_original = clean_text(_lookup(row, "title_original") or f"Truyện {index + 1}")
    title_localized = clean_text(_lookup(row, "title_localized"))
    raw_content = clean_text(_lookup(row, "raw_content"))
    content = clean_text(_lookup(row, "content") or raw_content)
    if not raw_content:
        raw_content = content
    source = clean_text(_lookup(row, "source") or "manual")
    source_url = str(_lookup(row, "source_url") or "").strip()
    language = str(_lookup(row, "language") or detect_language(raw_content)).lower()
    if language.startswith("zh"):
        language = "zh"
    elif language.startswith("en"):
        language = "en"
    else:
        language = "vi"
    record_id = str(_lookup(row, "id") or "").strip()
    if not record_id:
        record_id = _stable_id(source, source_url, title_original, raw_content)
    created = str(_lookup(row, "published_at") or "").strip()
    return {
        "id": record_id,
        "title_original": title_original,
        "title_localized": title_localized,
        "source": source,
        "source_url": source_url,
        "source_fingerprint": canonical_source_url(source_url),
        "author": clean_text(_lookup(row, "author")),
        "published_at": created,
        "language": language,
        "raw_content": raw_content,
        "content": content,
        "word_count": int(row.get("wordCount") or row.get("word_count") or
                          count_words(content, "vi" if title_localized else language)),
        "source_tags": _list_value(_lookup(row, "tags")),
        "notes": clean_text(_lookup(row, "notes")),
        "selected": bool(row.get("selected", False)),
        "status": str(row.get("status") or "raw"),
        "imported_at": datetime.now().isoformat(timespec="seconds"),
    }


def _json_rows(payload: Any) -> List[Dict[str, Any]]:
    if isinstance(payload, list):
        return [x for x in payload if isinstance(x, dict)]
    if isinstance(payload, dict):
        for key in ("stories", "items", "data", "results", "records"):
            value = payload.get(key)
            if isinstance(value, list):
                return [x for x in value if isinstance(x, dict)]
        return [payload]
    return []


def load_json(path: str, limit: int = 0) -> List[Dict[str, Any]]:
    with open(path, "r", encoding="utf-8-sig") as handle:
        rows = _json_rows(json.load(handle))
    if limit > 0:
        rows = rows[:limit]
    return [normalize_record(row, i) for i, row in enumerate(rows)]


def _sqlite_table(conn: sqlite3.Connection) -> str:
    tables = [row[0] for row in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
    ).fetchall()]
    aliases = {a.casefold() for key in ("content", "raw_content")
               for a in FIELD_ALIASES[key]}
    scored: List[Tuple[int, str]] = []
    for table in tables:
        quoted = table.replace('"', '""')
        cols = [str(row[1]) for row in conn.execute(
            f'PRAGMA table_info("{quoted}")').fetchall()]
        score = sum(1 for col in cols if col.casefold() in aliases)
        if score:
            score += 2 if table.casefold() in {"stories", "story", "articles"} else 0
            scored.append((score, table))
    if not scored:
        raise ValueError("SQLite không có bảng chứa cột nội dung truyện phù hợp.")
    return max(scored)[1]


def load_sqlite(path: str, limit: int = 0, table: str = "") -> List[Dict[str, Any]]:
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    try:
        chosen = table or _sqlite_table(conn)
        quoted = chosen.replace('"', '""')
        sql = f'SELECT * FROM "{quoted}"'
        params: Tuple[Any, ...] = ()
        if limit > 0:
            sql += " LIMIT ?"
            params = (limit,)
        rows = [dict(row) for row in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()
    return [normalize_record(row, i) for i, row in enumerate(rows)]


def load_source(path: str, limit: int = 0, table: str = "") -> List[Dict[str, Any]]:
    source = os.path.abspath(os.path.expandvars(str(path or "").strip().strip('"')))
    if not os.path.isfile(source):
        raise FileNotFoundError("Không tìm thấy file dữ liệu: " + source)
    ext = os.path.splitext(source)[1].lower()
    if ext in {".sqlite", ".sqlite3", ".db"}:
        return load_sqlite(source, limit=limit, table=table)
    if ext == ".json":
        return load_json(source, limit=limit)
    raise ValueError("Chỉ hỗ trợ file JSON, SQLite, SQLite3 hoặc DB.")


def _term_score(text: str, terms: Iterable[str]) -> int:
    folded = text.casefold()
    return sum(folded.count(term.casefold()) for term in terms)


def _rank_labels(text: str, groups: Dict[str, Sequence[str]], limit: int) -> List[str]:
    scores = [(name, _term_score(text, terms)) for name, terms in groups.items()]
    found = [name for name, score in sorted(scores, key=lambda item: item[1], reverse=True)
             if score > 0]
    return found[:limit]


def _headline_situation(record: Dict[str, Any], theme: str) -> str:
    localized = str(record.get("title_localized") or "").strip()
    language = str(record.get("language") or "").lower()
    # Chỉ tận dụng tiêu đề nguồn khi nó đã là tiếng Việt. Tiêu đề Trung/Anh
    # không được giả làm "Tiêu đề Việt" trong chế độ dự phòng offline.
    title = localized or (record.get("title_original") if language == "vi" else "") or ""
    title = re.sub(r"[|:—–-]+", " ", str(title)).strip()
    title = re.sub(r"\s+", " ", title)
    if 16 <= len(title) <= 62:
        return title.upper()
    situations = {
        "Tâm linh ông bà kể": "BÀ NGOẠI KỂ LẠI TIẾNG GÕ CỬA SAU ĐÊM GIỖ",
        "Mẹ chồng nàng dâu": "MẸ CHỒNG ĐUỔI CON DÂU KHỎI NHÀ GIỮA ĐÊM",
        "Con cái bất hiếu": "BA NGƯỜI CON TRANH NHÀ NHƯNG KHÔNG AI NUÔI MẸ",
        "Tình yêu xế chiều": "BÀ 68 TUỔI TÁI HÔN KHIẾN CẢ NHÀ PHẢN ĐỐI",
        "Tranh chấp thừa kế": "TỜ DI CHÚC XUẤT HIỆN NGAY TRƯỚC NGÀY CHIA ĐẤT",
        "Nghiệp báo ngoại tình": "NGƯỜI CHỒNG GIẤU BÍ MẬT SUỐT 20 NĂM",
        "Cô đơn tuổi già": "BỐ GIÀ BỊ BỎ LẠI TRONG CĂN NHÀ TRỐNG",
    }
    return situations.get(theme, "BÍ MẬT GIA ĐÌNH BẤT NGỜ BỊ PHƠI BÀY")


def _make_titles(record: Dict[str, Any], theme: str, emotion: str) -> List[str]:
    situation = _headline_situation(record, theme)
    spiritual = theme == "Tâm linh ông bà kể"
    hooks = ([
        "Nghe Mà Lạnh Gáy", "Chuyện Ông Bà Kể", "Nghe Mà Rùng Mình",
        "Chuyện Làng Quê Kỳ Bí", "Nghe Mà Thấm", "Nghe Mà Sốc",
        "Chuyện Cũ Bên Bếp Lửa", "Kể Chuyện Đêm Khuya",
    ] if spiritual else [
        "Nghe Mà Thấm", "Nghe THẤM Tận Xương", "Nghe Là Khóc",
        "Nghe Mà Nghẹn Lòng", "Nghe Mà Rơi Nước Mắt", "Nghe Sướng Lỗ Tai",
        "Nghe Mà Sốc", "Truyện Ngắn Tuổi Xế Chiều Cực Hay",
    ])
    tails = ([
        "Kể Chuyện Tâm Linh", "Kể Chuyện Đêm Khuya", "Kể Chuyện Làng Quê",
        "Kể Chuyện Tâm Linh", "Chuyện Ông Bà Kể", "Kể Chuyện Đêm Khuya",
        "Kể Chuyện Làng Quê", "Kể Chuyện Tâm Linh",
    ] if spiritual else [
        "Kể Chuyện Đêm Khuya", "Đọc Truyện Đêm Khuya", "Kể Chuyện Tuổi Già",
        "Kể Chuyện Làng Quê", "Kể Chuyện Đêm Khuya", "Kể Chuyện Tuổi Già",
        "Đọc Truyện Đêm Khuya", "Kể Chuyện Làng Quê",
    ])
    variations = [
        situation,
        situation.replace("BÍ MẬT", "SỰ THẬT"),
        situation.replace("NGƯỜI CHỒNG", "CHỒNG GIÀ"),
        situation.replace("CẢ NHÀ", "CÁC CON"),
        situation + " SAU BỮA CƠM GIỖ",
        situation.replace("GIỮA ĐÊM", "TRƯỚC MẶT HỌ HÀNG"),
        situation.replace("20 NĂM", "30 NĂM"),
        situation.replace("BẤT NGỜ", "SAU NGÀY GIỖ"),
    ]
    out: List[str] = []
    for hook, body, tail in zip(hooks, variations, tails):
        value = f"{hook} : {body} | {tail}"
        if len(value) > 100:
            room = max(18, 100 - len(hook) - len(tail) - 7)
            body = body[:room].rsplit(" ", 1)[0]
            value = f"{hook} : {body} | {tail}"
        if value not in out:
            out.append(value)
    while len(out) < 8:
        out.append(f"Nghe Mà Thấm : {situation[:48]} | Kể Chuyện Tuổi Già")
    return out[:8]


def _fallback_rewrite_material(theme: str, archetypes: Sequence[str],
                               situation: str) -> Dict[str, Any]:
    """Hồ sơ an toàn khi chưa gọi được AI: không chép một câu nào từ nguồn."""
    lead = archetypes[0] if archetypes else "người thân trong gia đình"
    scenes = [
        f"Biến cố mở màn cụ thể quanh {situation.lower()}, xảy ra trước mặt người thân.",
        f"Một dấu hiệu nhỏ cho thấy {lead.lower()} đang che giấu hoặc chịu đựng điều gì đó.",
        "Một lời nói dối để giữ hòa khí khiến một người bị tổn thương thật.",
        f"Áp lực tiền bạc hoặc lời dị nghị trong xóm làm mâu thuẫn {theme.lower()} bùng lên.",
        "Một vật chứng đời thường xuất hiện, đảo ngược cách mọi người nhìn sự việc.",
    ]
    twists = [
        "Người tưởng là có lỗi lại đang âm thầm bảo vệ một người khác.",
        "Quyền lợi tiền bạc chỉ là bề mặt; món nợ tình thân cũ mới là nguyên nhân chính.",
    ]
    brief = (
        "MỤC TIÊU: Viết một truyện gia đình Việt Nam hoàn toàn mới, không dịch và "
        "không kể lại truyện nguồn.\n"
        f"TIỀN ĐỀ MỚI: Khai thác chủ đề {theme.lower()} qua một gia đình ở làng quê "
        f"Nam Bộ. Mở truyện bằng sự việc: {situation.lower()}.\n"
        "NHỊP KỊCH TÍNH: " + " ".join(f"({i + 1}) {scene}" for i, scene in enumerate(scenes)) + "\n"
        "CHẤT LIỆU CÚ LẬT: " + " ".join(twists) + "\n"
        "RÀNG BUỘC SÁNG TÁC: Đổi toàn bộ tên, địa danh, quan hệ, số liệu, vật chứng, "
        "thứ tự sự việc và lời văn; thêm nguyên nhân-hậu quả mới; kết theo lẽ nhân quả. "
        "Không dùng tên riêng hoặc câu chữ của nguồn."
    )
    if theme == "Tâm linh ông bà kể":
        brief += (
            "\nQUY ƯỚC TÂM LINH: Kể qua lời hồi tưởng của ông/bà như chuyện truyền "
            "miệng bên bếp lửa; giữ không khí làng quê và sự mơ hồ. Không khẳng "
            "định hiện tượng siêu nhiên là thật, không hướng dẫn cúng bái/chữa bệnh "
            "mê tín; kết lại bằng tình thân hoặc bài học sống."
        )
    return {"main_hook": scenes[0], "high_tension_scenes": scenes,
            "plot_twists": twists,
            "must_change": ["tên", "địa danh", "quan hệ", "số liệu",
                            "vật chứng", "diễn biến", "lời văn"],
            "rewrite_brief": brief}


def _make_descriptions(record: Dict[str, Any], theme: str, emotion: str) -> List[str]:
    source_note = "Câu chuyện được sáng tạo lại từ một mô-típ gia đình, toàn bộ nhân vật và bối cảnh đều hư cấu."
    return [
        f"Một biến cố về {theme.lower()} đã đẩy cả gia đình đến lúc phải đối diện sự thật. "
        f"Câu chuyện mang màu sắc {emotion.lower()}, có nhiều lớp mâu thuẫn và những lựa chọn khiến người nghe phải suy ngẫm. "
        f"{source_note} Cô chú, anh chị hãy để lại cảm nhận và đăng ký để đồng hành cùng những câu chuyện mới. #kechuyen #tuoigia #giadinh",
        f"Khi tình thân bị thử thách bởi {theme.lower()}, ai mới là người giữ được nghĩa tình? "
        "Bản kể tập trung vào xung đột đời thường, cách đối nhân xử thế và bài học dành cho mỗi gia đình Việt. "
        f"{source_note} Mời mọi người nghe đến cuối và chia sẻ góc nhìn của mình. #truyendemkhuya #langque #chuyengiadinh",
        f"Từ một tình huống {emotion.lower()}, câu chuyện mở ra những bí mật, hiểu lầm và món nợ tình thân kéo dài nhiều năm. "
        "Nội dung đã được biên tập thành truyện Việt mới, không phải bản dịch nguyên văn. "
        "Nếu thấy ý nghĩa, cô chú anh chị nhớ bật chuông và để lại bình luận. #goctruyen #kechuyentuoiGia #nhanqua",
    ]


def heuristic_analysis(record: Dict[str, Any]) -> Dict[str, Any]:
    text = "\n".join([
        str(record.get("title_original") or ""),
        str(record.get("title_localized") or ""),
        str(record.get("content") or record.get("raw_content") or ""),
    ])
    themes = _rank_labels(text, THEME_TERMS, 4) or ["Gia đình và lòng tham"]
    emotions = _rank_labels(text, EMOTION_TERMS, 2) or ["Nghẹn lòng"]
    archetypes = _rank_labels(text, ARCHETYPE_TERMS, 6) or ["Người thân trong gia đình"]
    hook_hits = _term_score(text, HOOK_TERMS)
    twist_hits = _term_score(text, TWIST_TERMS)
    has_number = bool(re.search(r"\b\d{1,3}\b", text))
    has_quote = bool(re.search(r"[\"“”']", text))
    word_count = int(record.get("word_count") or count_words(text))
    hook = min(10, max(1, 4 + min(4, hook_hits) + int(has_number) + int(has_quote)))
    twist = min(10, max(1, 3 + min(5, twist_hits) + int(has_number)))
    duration = round(max(1, word_count / 150.0), 1)
    ease = "Cao" if 1200 <= word_count <= 15000 else "Trung" if word_count <= 24000 else "Thấp"
    priority = round(hook * .45 + twist * .40 + ({"Cao": 10, "Trung": 6, "Thấp": 3}[ease]) * .15, 1)
    recommendation = "Nên làm" if priority >= 7.2 else "Chờ đợi" if priority >= 5.5 else "Không nên làm"
    theme = themes[0]
    emotion = emotions[0]
    titles = _make_titles(record, theme, emotion)
    situation = _headline_situation(record, theme)
    material = _fallback_rewrite_material(theme, archetypes, situation)
    descriptions = _make_descriptions(record, theme, emotion)
    thumbnails = [
        {"description": f"Cận cảnh hai thế hệ đối đầu trong căn nhà làng quê, chủ đề {theme.lower()}",
         "colors": "đỏ sẫm, vàng ấm, đen", "text": "SỰ THẬT LỘ RA"},
        {"description": "Người mẹ già ngồi bên mâm cơm nguội, phía sau là các con đang tranh cãi",
         "colors": "xanh đêm, cam, trắng", "text": "KHÔNG AI NUÔI MẸ"},
        {"description": "Một tờ giấy quan trọng trên bàn thờ, mọi người bàng hoàng nhìn nhau",
         "colors": "nâu gỗ, đỏ, vàng", "text": "TỜ GIẤY CUỐI CÙNG"},
    ]
    keywords = list(dict.fromkeys(themes + [emotion, "chuyện gia đình", "tuổi già", "nhân quả"]))[:8]
    tags = ["kể chuyện đêm khuya", "chuyện tuổi già", "chuyện làng quê"] + themes[:3]
    if theme == "Tâm linh ông bà kể":
        tags = ["chuyện tâm linh", "chuyện ông bà kể", "chuyện làng quê",
                "kể chuyện đêm khuya"] + themes[:3]
    series = f"Series {theme} · {archetypes[0]}"
    outline = (
        f"Phần 1 — Mở nút: Giới thiệu {archetypes[0].lower()} và biến cố về {theme.lower()}.\n"
        f"Phần 2 — Leo thang: Mâu thuẫn gia đình dồn dập, các nhân vật buộc phải chọn bên.\n"
        "Phần 3 — Cao trào: Một sự thật quan trọng xuất hiện; dừng trước kết cục để biên tập phát triển truyện mới."
    )
    return {
        "title_localized": situation.title(),
        "primary_genre": theme,
        "emotion": emotion,
        "hook_score": hook,
        "plot_twist_score": twist,
        "themes": themes,
        "archetypes": archetypes,
        "keywords": keywords,
        "tags": list(dict.fromkeys(tags)),
        "estimated_duration_minutes": duration,
        "series": series,
        "titles": titles,
        **material,
        "descriptions": descriptions,
        "thumbnails": thumbnails,
        "outline": outline,
        "main_characters": archetypes[:5],
        "production_ease": ease,
        "priority_score": priority,
        "recommendation": recommendation,
        "best_publish_time": "19:30–21:00, ưu tiên Thứ Năm hoặc Chủ Nhật",
        "analysis_provider": "heuristic",
        "analysis_error": "",
        "status": "analyzed",
    }


def _json_object_candidates(text: str) -> List[str]:
    """Tách các object JSON cân bằng, bỏ qua ngoặc nằm trong chuỗi.

    Gemini Web đôi khi thêm lời dẫn hoặc bọc kết quả trong Markdown. Cắt từ dấu
    ``{`` đầu tới dấu ``}`` cuối rất dễ nuốt hai khối khác nhau vào cùng một
    chuỗi và làm ``json.loads`` thất bại, nên ta quét từng object độc lập.
    """
    raw = str(text or "").lstrip("\ufeff").strip()
    candidates: List[str] = []
    start = -1
    depth = 0
    in_string = False
    escaped = False
    for index, char in enumerate(raw):
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            if depth == 0:
                start = index
            depth += 1
        elif char == "}" and depth:
            depth -= 1
            if depth == 0 and start >= 0:
                candidates.append(raw[start:index + 1])
                start = -1
    return candidates


def _remove_json_trailing_commas(text: str) -> str:
    """Bỏ dấu phẩy cuối object/array mà không sửa nội dung trong chuỗi."""
    out: List[str] = []
    in_string = False
    escaped = False
    index = 0
    while index < len(text):
        char = text[index]
        if in_string:
            out.append(char)
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            index += 1
            continue
        if char == '"':
            in_string = True
            out.append(char)
            index += 1
            continue
        if char == ",":
            lookahead = index + 1
            while lookahead < len(text) and text[lookahead].isspace():
                lookahead += 1
            if lookahead < len(text) and text[lookahead] in "}]":
                index += 1
                continue
        out.append(char)
        index += 1
    return "".join(out)


def _extract_json(text: str) -> Dict[str, Any]:
    raw = str(text or "").lstrip("\ufeff").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I)
    candidates = _json_object_candidates(raw)
    if not candidates:
        raise ValueError("AI không trả về JSON.")
    last_error: Optional[Exception] = None
    for candidate in candidates:
        for variant in (candidate, _remove_json_trailing_commas(candidate)):
            for strict in (True, False):
                try:
                    parsed = json.loads(variant, strict=strict)
                except (TypeError, ValueError) as exc:
                    last_error = exc
                    continue
                if isinstance(parsed, dict):
                    return parsed
                last_error = ValueError("JSON AI không phải object.")
    detail = str(last_error or "JSON không hợp lệ").splitlines()[0][:160]
    raise ValueError(f"AI trả JSON chưa hợp lệ: {detail}")


def _browser_json_correction_prompt(response: str) -> str:
    """Yêu cầu model tự định dạng lại câu vừa trả, không phân tích nguồn lần hai."""
    previous = str(response or "").strip()
    return (
        "Câu trả lời vừa rồi chưa đọc được bằng JSON. Hãy XUẤT LẠI TOÀN BỘ "
        "kết quả phân tích thành đúng MỘT JSON object theo schema tôi đã yêu cầu. "
        "Ký tự đầu tiên phải là { và ký tự cuối cùng phải là }. Không dùng khối "
        "```json, không lời dẫn, không chú thích, không dấu phẩy thừa. Mọi xuống "
        "dòng bên trong chuỗi phải viết thành \\n. Không phân tích lại và không bỏ trường.\n\n"
        "CÂU TRẢ LỜI CẦN ĐỊNH DẠNG LẠI:\n" + previous[:30000]
    )


def _save_invalid_ai_response(config: Dict[str, Any], provider: str,
                              model: str, response: str) -> str:
    """Giữ câu trả lời lỗi gần nhất để lần sau không phải đoán từ log."""
    cp = config.get("content_pipeline") if isinstance(
        config.get("content_pipeline"), dict) else {}
    configured = str(cp.get("output_dir") or DEFAULT_OUTPUT_DIR).strip()
    folder = os.path.abspath(
        configured if os.path.isabs(configured) else os.path.join(ROOT, configured))
    path = os.path.join(folder, "_phan_hoi_ai_json_loi.txt")
    try:
        os.makedirs(folder, exist_ok=True)
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(
                f"Thời gian: {datetime.now().isoformat(timespec='seconds')}\n"
                f"Provider: {provider}\nModel: {model}\n\n{str(response or '')}")
        return path
    except OSError:
        return ""


def _score(value: Any, fallback: int) -> int:
    try:
        return max(1, min(10, int(round(float(value)))))
    except (TypeError, ValueError):
        return fallback


def _normalize_ai_result(ai: Dict[str, Any], fallback: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(fallback)
    scalar = (
        "title_localized", "primary_genre", "emotion", "series", "outline",
        "main_hook", "rewrite_brief", "production_ease", "recommendation",
        "best_publish_time")
    for key in scalar:
        value = str(ai.get(key) or "").strip()
        if value:
            out[key] = value
    out["hook_score"] = _score(ai.get("hook_score"), fallback["hook_score"])
    out["plot_twist_score"] = _score(
        ai.get("plot_twist_score"), fallback["plot_twist_score"])
    for key, wanted in (("themes", 4), ("archetypes", 6), ("keywords", 10),
                        ("tags", 10), ("main_characters", 8),
                        ("high_tension_scenes", 7), ("plot_twists", 4),
                        ("must_change", 12)):
        values = _list_value(ai.get(key))
        if values:
            out[key] = values[:wanted]
    titles = _list_value(ai.get("titles"))
    if titles:
        out["titles"] = list(dict.fromkeys(titles + fallback["titles"]))[0:8]
    descriptions = _list_value(ai.get("descriptions"))
    if descriptions:
        out["descriptions"] = (descriptions + fallback["descriptions"])[0:3]
    thumbs = ai.get("thumbnails")
    if isinstance(thumbs, list):
        clean_thumbs = []
        for item in thumbs[:3]:
            if isinstance(item, dict):
                clean_thumbs.append({
                    "description": str(item.get("description") or "").strip(),
                    "colors": str(item.get("colors") or "").strip(),
                    "text": str(item.get("text") or "").strip(),
                })
            elif str(item).strip():
                clean_thumbs.append({"description": str(item).strip(),
                                     "colors": "", "text": ""})
        if clean_thumbs:
            out["thumbnails"] = (clean_thumbs + fallback["thumbnails"])[:3]
    hook, twist = out["hook_score"], out["plot_twist_score"]
    ease_value = {"Cao": 10, "Trung": 6, "Thấp": 3}.get(out["production_ease"], 6)
    out["priority_score"] = round(hook * .45 + twist * .40 + ease_value * .15, 1)
    out["analysis_provider"] = "ai"
    out["analysis_error"] = ""
    out["status"] = "analyzed"
    return out


_BROWSER_ANALYSIS_PROVIDERS = {"browser", "perplexity_browser"}


def _browser_profile_path(value: Any, fallback: str) -> str:
    raw = os.path.expandvars(os.path.expanduser(str(value or fallback).strip().strip('"')))
    return os.path.abspath(raw if os.path.isabs(raw) else os.path.join(ROOT, raw))


def _perplexity_tool_settings(config: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    section = config.get("tao_kich_ban") if isinstance(
        config.get("tao_kich_ban"), dict) else {}
    tool_dir = _browser_profile_path(
        section.get("tool_dir"), os.path.join(os.pardir, "Tạo kịch bản"))
    config_path = os.path.join(tool_dir, "config.json")
    settings: Dict[str, Any] = {}
    if os.path.isfile(config_path):
        with open(config_path, "r", encoding="utf-8") as handle:
            loaded = json.load(handle)
        if isinstance(loaded, dict):
            settings.update(loaded)
    cp = config.get("content_pipeline") if isinstance(
        config.get("content_pipeline"), dict) else {}
    settings.setdefault("site_url", "https://www.perplexity.ai/")
    settings.setdefault("browser_channel", "chrome")
    settings.setdefault("preferred_model", str(
        cp.get("perplexity_model") or "Claude Sonnet 5"))
    settings.setdefault("model_reasoning", True)
    settings.setdefault("answer_timeout_sec", 420)
    settings.setdefault("first_token_timeout_sec", 180)
    settings.setdefault("empty_fail_seconds", 90)
    settings.setdefault("stable_seconds", 6.0)
    settings.setdefault("poll_interval_sec", 1.5)
    profile = str(settings.get("user_data_dir") or "browser_profile_perplexity")
    settings["user_data_dir"] = os.path.abspath(
        profile if os.path.isabs(profile) else os.path.join(tool_dir, profile))
    return tool_dir, settings


@contextmanager
def _browser_analysis_session(config: Dict[str, Any], provider: str):
    """Mở một phiên web đã đăng nhập và trả hàm hỏi; không dùng API/key."""
    from . import translate

    provider_name = str(provider or "browser").strip().lower()
    if provider_name == "browser":
        tr = config.get("translation") if isinstance(
            config.get("translation"), dict) else {}
        profile = _browser_profile_path(
            tr.get("browser_profile"), "browser_profile")
        channel = str(tr.get("browser_channel") or "msedge")
        url = str(tr.get("browser_url") or "https://gemini.google.com/app")
        wait_reply = max(60, int(tr.get("wait_reply") or 240))
        with translate.phien_gemini_trinh_duyet(
                profile, channel=channel, url=url, wait_reply=wait_reply) as ask:
            yield ask
        return

    tool_dir, settings = _perplexity_tool_settings(config)
    driver_path = os.path.join(tool_dir, "app", "perplexity.py")
    if not os.path.isfile(driver_path):
        raise RuntimeError(
            "Không thấy bộ điều khiển Perplexity tại %s" % driver_path)
    module_name = "_autodub_perplexity_browser"
    module = sys.modules.get(module_name)
    if module is None:
        spec = importlib.util.spec_from_file_location(module_name, driver_path)
        if spec is None or spec.loader is None:
            raise RuntimeError("Không nạp được bộ điều khiển Perplexity.")
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
    driver = module.PerplexityDriver(settings, translate.log)
    try:
        driver.start()
        driver.goto_home()
        if not driver.is_logged_in():
            translate.log(
                "Perplexity chưa nhận phiên Pro; hãy đăng nhập trong cửa sổ Chrome vừa mở. "
                "Chương trình sẽ tự tiếp tục sau khi đăng nhập.",
                "warn")
            deadline = time.time() + 300
            while time.time() < deadline and not driver.is_logged_in():
                driver._sleep(3)
            if not driver.is_logged_in():
                raise RuntimeError(
                    "Sau 5 phút vẫn chưa thấy phiên đăng nhập Perplexity Pro; "
                    "không dùng tài khoản khách thay cho Claude đã chọn.")
        driver.wait_for_input(timeout=180)
        driver.ensure_model()

        def ask(prompt: str) -> str:
            return driver.ask(prompt, new_thread=True,
                              label="phân tích ý tưởng bằng Claude/Perplexity")

        yield ask
    finally:
        driver.close()


def _provider_params(config: Dict[str, Any], requested: str = "auto"
                     ) -> Tuple[str, str, str, Optional[str], int]:
    from . import translate
    tr = config.get("translation") if isinstance(config.get("translation"), dict) else {}
    cp = config.get("content_pipeline") if isinstance(config.get("content_pipeline"), dict) else {}
    provider = str(requested or cp.get("provider") or "auto").strip().lower()
    if provider == "auto":
        provider = str(tr.get("provider") or "browser").strip().lower()
    if provider == "browser":
        return "browser", "", "Gemini Web", None, max(
            60, int(tr.get("wait_reply") or 240))
    if provider == "perplexity_browser":
        _tool_dir, settings = _perplexity_tool_settings(config)
        model = str(settings.get("preferred_model") or "Claude trên Perplexity")
        return "perplexity_browser", "", model, None, int(
            settings.get("answer_timeout_sec") or 420)
    if provider in {"", "none", "heuristic", "offline"}:
        return provider or "heuristic", "", "", None, 420
    key, model, base_url, timeout = translate.api_params_for_provider(tr, provider)
    return provider, str(key or ""), str(model or ""), base_url, int(timeout or 420)


def _analysis_error_message(provider: str, model: str, exc: Exception) -> str:
    """Bien loi API thanh thong bao co huong sua, khong lam lo credential."""
    raw = str(exc or "Lỗi không xác định").strip()
    lowered = raw.casefold()
    label = str(provider or "AI").upper()
    model_note = f" / {model}" if model else ""
    if any(token in lowered for token in ("401", "unauthorized", "invalid api key",
                                           "incorrect api key")):
        advice = "API key sai hoặc hết hiệu lực; mở Cài đặt → API, nhập lại key rồi Kiểm tra kết nối."
    elif any(token in lowered for token in ("403", "forbidden", "permission")):
        advice = "Tài khoản/key chưa được cấp quyền dùng model này; đổi model hoặc key."
    elif any(token in lowered for token in ("429", "rate limit", "quota", "too many")):
        advice = "API hết hạn mức hoặc bị giới hạn tốc độ; chờ rồi thử lại, hoặc đổi dịch vụ AI."
    elif any(token in lowered for token in ("404", "model_not_found", "model not found")):
        advice = "Không tìm thấy model; kiểm tra lại tên model trong Cài đặt → API."
    elif any(token in lowered for token in ("timeout", "timed out", "time out")):
        advice = "API chờ quá lâu; kiểm tra mạng, giảm số tác vụ đồng thời hoặc đổi dịch vụ AI."
    elif any(token in lowered for token in ("connection", "dns", "ssl", "network")):
        advice = "Không kết nối được dịch vụ AI; kiểm tra mạng, proxy/tường lửa rồi thử lại."
    elif "json" in lowered:
        advice = "AI trả sai định dạng; thử lại hoặc đổi model ổn định hơn."
    else:
        advice = "Mở Cài đặt → API, Kiểm tra kết nối; nếu vẫn lỗi hãy đổi model/dịch vụ AI."
    return f"{label}{model_note} lỗi: {raw[:220]}. Cách sửa: {advice}"


def analyze_record(record: Dict[str, Any], config: Optional[Dict[str, Any]] = None,
                   use_ai: bool = True, provider: str = "auto",
                   progress: Optional[Callable[[str, str, float], None]] = None,
                   browser_ask: Optional[Callable[[str], str]] = None
                   ) -> Dict[str, Any]:
    def notify(stage: str, message: str, pct: float) -> None:
        if progress:
            progress(stage, message, max(0.0, min(100.0, float(pct))))

    base = heuristic_analysis(record)
    merged = dict(record)
    if not use_ai:
        notify("offline", "Đang phân tích offline (không gọi API AI).", 40)
        merged.update(base)
        notify("done", "Đã phân tích offline.", 100)
        return merged
    cfg = config or {}
    provider_name, api_key, model, base_url, timeout = _provider_params(cfg, provider)
    is_browser = provider_name in _BROWSER_ANALYSIS_PROVIDERS
    if not api_key and not is_browser:
        if provider_name == "nvidia":
            base["analysis_error"] = (
                "NVIDIA chưa có API key trong Cài đặt > API; "
                "đã dùng phân tích offline.")
        else:
            base["analysis_error"] = (
                "Không có API key cho provider %s; đã dùng phân tích offline."
                % provider_name)
        merged.update(base)
        notify("warning", base["analysis_error"], 100)
        return merged
    content = str(record.get("content") or record.get("raw_content") or "")
    if len(content) <= 24000:
        excerpt = content
    else:
        middle = max(8000, len(content) // 2 - 4000)
        excerpt = (content[:8000] + "\n\n[ĐOẠN GIỮA NGUỒN]\n" +
                   content[middle:middle + 8000] +
                   "\n\n[ĐOẠN CUỐI NGUỒN]\n" + content[-8000:])
    prompt = (
        ANALYSIS_PROMPT + "\n\nDỮ LIỆU NGUỒN:\n"
        + json.dumps({
            "title": record.get("title_localized") or record.get("title_original"),
            "source": record.get("source"),
            "language": record.get("language"),
            "word_count": record.get("word_count"),
            "content_excerpt_begin_middle_end": excerpt,
        }, ensure_ascii=False)
    )
    try:
        from . import translate
        ai_label = str(provider_name or "AI").upper()
        notify("sending", f"Đang gửi nội dung tới {ai_label} · model {model or 'mặc định'}…", 20)
        notify("waiting", f"Đang chờ {ai_label} trả hook, plot twist và 8 tiêu đề…", 45)
        heartbeat_stop = threading.Event()
        wait_started = time.monotonic()

        def heartbeat() -> None:
            while not heartbeat_stop.wait(20):
                waited = int(time.monotonic() - wait_started)
                notify("waiting", (f"Vẫn đang chờ {ai_label} · {model or 'model mặc định'} "
                                   f"phản hồi ({waited} giây)…"), 45)

        if progress:
            threading.Thread(target=heartbeat, daemon=True).start()
        try:
            if is_browser:
                if browser_ask is not None:
                    response = browser_ask(prompt)
                else:
                    with _browser_analysis_session(cfg, provider_name) as ask:
                        response = ask(prompt)
            else:
                response = translate._api_call(
                    prompt, api_key, model, 0.25, provider=provider_name,
                    api_base_url=base_url,
                    # Phân tích một ý tưởng luôn có kết quả heuristic dự phòng. Không
                    # retry 5 lần khi NVIDIA giữ request đến lúc gateway 504: trước
                    # đây một bài có thể làm giao diện chờ hơn 25 phút.
                    api_timeout=min(int(timeout or 420), 180),
                    api_retries=1 if provider_name == "nvidia" else None)
        finally:
            heartbeat_stop.set()
        notify("parsing", f"{ai_label} đã trả lời; đang kiểm tra JSON và lưu kết quả…", 85)
        try:
            parsed_response = _extract_json(response)
        except ValueError as first_parse_error:
            cp = cfg.get("content_pipeline") if isinstance(
                cfg.get("content_pipeline"), dict) else {}
            try:
                configured_retries = int(cp.get("browser_json_retries", 1) or 0)
            except (TypeError, ValueError):
                configured_retries = 1
            json_retries = max(0, min(2, configured_retries))
            if not is_browser or json_retries <= 0:
                raise
            last_parse_error: Exception = first_parse_error
            for retry_index in range(json_retries):
                notify(
                    "parsing",
                    f"{ai_label} trả chưa đúng JSON; đang yêu cầu xuất lại "
                    f"({retry_index + 1}/{json_retries})…",
                    88,
                )
                if browser_ask is not None:
                    response = browser_ask(_browser_json_correction_prompt(response))
                else:
                    # Phiên đầu đã đóng ở nhánh gọi đơn lẻ, nên phiên mới phải
                    # nhận lại cả đề bài thay vì câu ngắn "schema đã yêu cầu".
                    correction = (
                        prompt + "\n\nYÊU CẦU ĐỊNH DẠNG BẮT BUỘC: ký tự đầu "
                        "tiên của câu trả lời là {, ký tự cuối là }; chỉ một "
                        "JSON object hợp lệ, không Markdown hay lời dẫn."
                    )
                    with _browser_analysis_session(cfg, provider_name) as ask:
                        response = ask(correction)
                try:
                    parsed_response = _extract_json(response)
                    break
                except ValueError as retry_error:
                    last_parse_error = retry_error
            else:
                debug_path = _save_invalid_ai_response(
                    cfg, provider_name, model, response)
                suffix = f" Phản hồi lỗi đã lưu tại: {debug_path}" if debug_path else ""
                message = str(last_parse_error).rstrip(".") + "."
                raise ValueError(f"{message}{suffix}") from last_parse_error
        analysis = _normalize_ai_result(parsed_response, base)
        analysis["analysis_provider"] = f"{provider_name}:{model}"
        notify("done", f"{ai_label} · {model or 'model mặc định'} phân tích thành công.", 100)
    except Exception as exc:
        analysis = base
        analysis["analysis_error"] = _analysis_error_message(
            provider_name, model, exc) + " Đã dùng kết quả offline thay thế."
        analysis["analysis_provider"] = "heuristic-fallback"
        notify("error", analysis["analysis_error"], 100)
    merged.update(analysis)
    return merged


async def analyze_many(records: Sequence[Dict[str, Any]], config: Dict[str, Any],
                       use_ai: bool = True, provider: str = "auto",
                       concurrency: int = 3,
                       progress: Optional[Callable[[int, int, Dict[str, Any]], None]] = None,
                       activity: Optional[Callable[[int, int, str, str, float], None]] = None
                       ) -> List[Dict[str, Any]]:
    resolved_provider = _provider_params(config, provider)[0]
    if use_ai and resolved_provider in _BROWSER_ANALYSIS_PROVIDERS:
        def browser_batch() -> List[Dict[str, Any]]:
            results: List[Dict[str, Any]] = []
            try:
                with _browser_analysis_session(config, resolved_provider) as ask:
                    for index, source in enumerate(records):
                        def report(stage: str, message: str, pct: float,
                                   item_index: int = index) -> None:
                            if activity:
                                activity(item_index + 1, len(records), stage, message, pct)

                        result = analyze_record(
                            dict(source), config, True, resolved_provider, report,
                            browser_ask=ask)
                        results.append(result)
                        if progress:
                            progress(len(results), len(records), result)
                return results
            except Exception as exc:
                # Không mở được profile/Playwright vẫn phải giữ đủ kết quả local.
                failed: List[Dict[str, Any]] = []
                for index, source in enumerate(records):
                    result = analyze_record(
                        dict(source), config, use_ai=False, provider="heuristic")
                    result["analysis_provider"] = "heuristic-fallback"
                    result["analysis_error"] = _analysis_error_message(
                        resolved_provider, _provider_params(config, resolved_provider)[2],
                        exc) + " Đã dùng kết quả offline thay thế."
                    failed.append(result)
                    if activity:
                        activity(index + 1, len(records), "error",
                                 result["analysis_error"], 100)
                    if progress:
                        progress(index + 1, len(records), result)
                return failed

        return await asyncio.to_thread(browser_batch)

    sem = asyncio.Semaphore(max(1, min(8, int(concurrency or 3))))
    done = 0
    lock = asyncio.Lock()

    async def one(index: int, record: Dict[str, Any]) -> Dict[str, Any]:
        nonlocal done
        def report(stage: str, message: str, pct: float) -> None:
            if activity:
                activity(index + 1, len(records), stage, message, pct)

        async with sem:
            result = await asyncio.to_thread(
                analyze_record, record, config, use_ai, provider, report)
        async with lock:
            done += 1
            if progress:
                progress(done, len(records), result)
        return result

    return list(await asyncio.gather(*(
        one(index, dict(record)) for index, record in enumerate(records))))


def record_was_used(record: Dict[str, Any]) -> bool:
    if not isinstance(record, dict):
        return False
    try:
        usage_count = int(record.get("usage_count") or 0)
    except (TypeError, ValueError):
        usage_count = 0
    if str(record.get("used_at") or "").strip() or usage_count > 0:
        return True
    status = str(record.get("status") or "").strip().casefold()
    # "Đang viết" có thể là trạng thái nhập sẵn từ kho cũ và không chứng minh
    # người dùng đã thực sự sản xuất. Chỉ khóa khi có dấu thời gian/lượt dùng,
    # hoặc trạng thái hoàn tất rõ ràng.
    return status in {"đã dùng", "used", "rendered", "published"}


class ContentStore:
    def __init__(self, path: str = DEFAULT_DB):
        self.path = os.path.abspath(path)
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        self._init()

    def connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path, timeout=30)
        conn.row_factory = sqlite3.Row
        return conn

    def _init(self) -> None:
        conn = self.connect()
        try:
            with conn:
                conn.execute("""
                CREATE TABLE IF NOT EXISTS content_ideas (
                    id TEXT PRIMARY KEY,
                    payload TEXT NOT NULL,
                    selected INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """)
                conn.execute("CREATE INDEX IF NOT EXISTS idx_content_selected "
                             "ON content_ideas(selected, updated_at)")
        finally:
            conn.close()

    def upsert(self, records: Sequence[Dict[str, Any]], overwrite: bool = False) -> int:
        now = datetime.now().isoformat(timespec="seconds")
        count = 0
        conn = self.connect()
        try:
            with conn:
                for record in records:
                    item = dict(record)
                    record_id = str(item.get("id") or "").strip()
                    if not record_id:
                        continue
                    old = conn.execute(
                        "SELECT payload, selected, created_at FROM content_ideas WHERE id=?",
                        (record_id,)).fetchone()
                    if old:
                        previous = json.loads(old["payload"])
                        if overwrite:
                            # Người dùng chủ động Phân tích lại: thay kết quả AI cũ,
                            # nhưng trạng thái tick vẫn do cột selected quản lý.
                            previous.update(item)
                        else:
                            # Nhập lại cùng kho nguồn không được xoá phần đã phân
                            # tích/chỉnh tay; chỉ bổ sung trường còn thiếu.
                            for key, value in item.items():
                                if key not in previous or previous.get(key) in (None, "", []):
                                    previous[key] = value
                        item = previous
                        selected = int(old["selected"])
                        created = str(old["created_at"])
                    else:
                        selected = int(bool(item.get("selected")))
                        created = now
                    item["selected"] = bool(selected)
                    conn.execute("""
                        INSERT INTO content_ideas(id,payload,selected,created_at,updated_at)
                        VALUES(?,?,?,?,?)
                        ON CONFLICT(id) DO UPDATE SET payload=excluded.payload,
                          selected=excluded.selected, updated_at=excluded.updated_at
                    """, (record_id, json.dumps(item, ensure_ascii=False), selected,
                          created, now))
                    count += 1
        finally:
            conn.close()
        return count

    def list(self, search: str = "", selected_only: bool = False,
             limit: int = 1000) -> List[Dict[str, Any]]:
        clauses, params = [], []
        if selected_only:
            clauses.append("selected=1")
        sql = "SELECT payload, selected FROM content_ideas"
        if clauses:
            sql += " WHERE " + " AND ".join(clauses)
        sql += " ORDER BY updated_at DESC LIMIT ?"
        params.append(max(1, min(10000, int(limit or 1000))))
        conn = self.connect()
        try:
            rows = conn.execute(sql, params).fetchall()
        finally:
            conn.close()
        items = []
        needle = str(search or "").casefold().strip()
        for row in rows:
            item = json.loads(row["payload"])
            item["selected"] = bool(row["selected"])
            if needle:
                haystack = " ".join(str(item.get(k) or "") for k in (
                    "title_original", "title_localized", "source", "primary_genre",
                    "emotion", "series")).casefold()
                if needle not in haystack:
                    continue
            items.append(item)
        return items

    def get(self, record_id: str) -> Optional[Dict[str, Any]]:
        conn = self.connect()
        try:
            row = conn.execute("SELECT payload,selected FROM content_ideas WHERE id=?",
                               (str(record_id),)).fetchone()
        finally:
            conn.close()
        if not row:
            return None
        item = json.loads(row["payload"])
        item["selected"] = bool(row["selected"])
        return item

    def source_history(self, urls: Sequence[str] = ()) -> Dict[str, Dict[str, Any]]:
        """Lập chỉ mục URL đã tải/đã dùng để chặn sản xuất trùng chuyện."""
        wanted = {canonical_source_url(x) for x in urls if canonical_source_url(x)}
        rows = self.list(limit=10000)
        history: Dict[str, Dict[str, Any]] = {}
        for item in rows:
            key = str(item.get("source_fingerprint") or
                      canonical_source_url(item.get("source_url") or ""))
            if not key or (wanted and key not in wanted):
                continue
            previous = history.get(key)
            used = record_was_used(item)
            value = {
                "id": item.get("id"),
                "title": item.get("title_localized") or item.get("title_original") or "",
                "source_url": item.get("source_url") or "",
                "used": used,
                "used_at": item.get("used_at") or "",
                "usage_count": (int(item.get("usage_count") or 0)
                                if str(item.get("usage_count") or "").isdigit()
                                else (1 if used else 0)),
                "status": item.get("status") or "",
            }
            if previous is None or (used and not previous.get("used")):
                history[key] = value
        return history

    def patch(self, record_id: str, values: Dict[str, Any]) -> Dict[str, Any]:
        item = self.get(record_id)
        if not item:
            raise KeyError("Không tìm thấy ý tưởng.")
        allowed = {
            "title_localized", "primary_genre", "emotion", "hook_score",
            "plot_twist_score", "themes", "archetypes", "keywords", "tags",
            "series", "titles", "descriptions", "thumbnails", "outline",
            "main_hook", "high_tension_scenes", "plot_twists", "must_change",
            "rewrite_brief",
            "main_characters", "production_ease", "recommendation",
            "best_publish_time", "notes", "status", "content", "selected",
            "used_at", "usage_count", "last_used_title",
        }
        for key, value in values.items():
            if key in allowed:
                item[key] = value
        self.upsert([item], overwrite=True)
        if "selected" in values:
            self.set_selected([record_id], bool(values["selected"]))
        return self.get(record_id) or item

    def set_selected(self, ids: Sequence[str], selected: bool) -> int:
        clean_ids = [str(x) for x in ids if str(x).strip()]
        if not clean_ids:
            return 0
        marks = ",".join("?" for _ in clean_ids)
        conn = self.connect()
        try:
            with conn:
                rows = conn.execute(
                    f"SELECT id,payload FROM content_ideas WHERE id IN ({marks})",
                    clean_ids).fetchall()
                for row in rows:
                    item = json.loads(row["payload"])
                    item["selected"] = bool(selected)
                    conn.execute(
                        "UPDATE content_ideas SET selected=?,payload=?,updated_at=? WHERE id=?",
                        (int(selected), json.dumps(item, ensure_ascii=False),
                         datetime.now().isoformat(timespec="seconds"), row["id"]))
        finally:
            conn.close()
        return len(rows)

    def delete_many(self, ids: Sequence[str], protect_used: bool = True
                    ) -> Dict[str, List[str]]:
        """Xóa nhiều mục; mặc định không cho xóa dấu vết chuyện đã dùng."""
        clean_ids = list(dict.fromkeys(
            str(value).strip() for value in ids if str(value).strip()))
        result = {"deleted": [], "protected": [], "not_found": []}
        if not clean_ids:
            return result
        marks = ",".join("?" for _ in clean_ids)
        conn = self.connect()
        try:
            with conn:
                rows = conn.execute(
                    f"SELECT id,payload FROM content_ideas WHERE id IN ({marks})",
                    clean_ids).fetchall()
                found = {str(row["id"]): row for row in rows}
                for record_id in clean_ids:
                    row = found.get(record_id)
                    if row is None:
                        result["not_found"].append(record_id)
                        continue
                    item = json.loads(row["payload"])
                    if protect_used and record_was_used(item):
                        result["protected"].append(record_id)
                        continue
                    conn.execute("DELETE FROM content_ideas WHERE id=?", (record_id,))
                    result["deleted"].append(record_id)
        finally:
            conn.close()
        return result


def public_record(record: Dict[str, Any], include_content: bool = False) -> Dict[str, Any]:
    item = dict(record)
    if not include_content:
        item["content_preview"] = str(item.get("content") or item.get("raw_content") or "")[:600]
        item.pop("content", None)
        item.pop("raw_content", None)
    return item


def _thumb_text(item: Any) -> str:
    if isinstance(item, dict):
        return "Hình: %s | Màu: %s | Chữ: %s" % (
            item.get("description", ""), item.get("colors", ""), item.get("text", ""))
    return str(item or "")


CONTENT_CALENDAR_HEADERS = [
    "STT", "Ngày tạo", "Thứ", "Ca đăng", "Nhóm chủ đề",
    "Tiêu đề đề xuất", "Thời lượng mục tiêu", "Nguồn kịch bản", "Trạng thái",
]


def _calendar_topic(record: Dict[str, Any]) -> str:
    themes = record.get("themes") or []
    if isinstance(themes, str):
        themes = [x.strip() for x in re.split(r"[,;\n]+", themes) if x.strip()]
    return str(record.get("primary_genre") or (themes[0] if themes else "") or
               record.get("series") or "").strip()


def _calendar_title(record: Dict[str, Any], final_title: str = "") -> str:
    titles = list(record.get("titles") or [])
    return str(final_title or record.get("last_used_title") or
               (titles[0] if titles else "") or record.get("title_localized") or
               record.get("title_original") or "").strip()


def _calendar_status(record: Dict[str, Any], default: str = "Chưa làm") -> str:
    raw = str(record.get("status") or "").strip()
    mapping = {
        "new": "Chưa làm", "downloaded": "Chưa làm", "analyzed": "Chưa làm",
        "pending": "Chưa làm", "used": "Đang viết", "writing": "Đang viết",
        "rendering": "Đang dựng", "rendered": "Đã xuất video",
        "published": "Đã đăng", "skipped": "Tạm hoãn",
    }
    return mapping.get(raw.casefold(), raw or default)


def _style_content_calendar(sheet, max_row: int) -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    navy, cyan = "0B1626", "14D8D4"
    header_fill = PatternFill("solid", fgColor=navy)
    thin = Side(style="thin", color="D7E1EA")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=cyan))
    sheet.row_dimensions[1].height = 34
    widths = (8, 14, 15, 25, 27, 68, 22, 26, 22, 22, 48)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=max(2, max_row), min_col=1, max_col=9):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 4)
            cell.border = Border(bottom=thin)
        row[1].number_format = "dd/mm/yyyy"
        for col in (4, 5, 6, 7, 8, 9):
            row[col - 1].fill = PatternFill("solid", fgColor="FFF7CC")
        sheet.row_dimensions[row[0].row].height = 38
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(1, max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["J"].hidden = True
    sheet.column_dimensions["K"].hidden = True
    if max_row >= 2:
        if not sheet.data_validations.dataValidation:
            dv = DataValidation(
                type="list",
                formula1='"Chưa làm,Đang viết,Đang dựng,Đã xuất video,Đã đăng,Tạm hoãn"')
            sheet.add_data_validation(dv)
            dv.add(f"I2:I{max_row}")
        else:
            sheet.data_validations.dataValidation[0].sqref = f"I2:I{max_row}"
        if not len(sheet.conditional_formatting):
            sheet.conditional_formatting.add(
                f"I2:I{max_row}", FormulaRule(formula=['$I2="Đã đăng"'],
                                                fill=PatternFill("solid", fgColor="DCFCE7")))
            sheet.conditional_formatting.add(
                f"I2:I{max_row}", FormulaRule(formula=['$I2="Tạm hoãn"'],
                                                fill=PatternFill("solid", fgColor="FEE2E2")))
    tables = list(sheet.tables.values())
    if tables:
        tables[0].ref = f"A1:I{max(2, max_row)}"
    elif max_row >= 2:
        table = Table(displayName="ContentCalendar", ref=f"A1:I{max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def append_content_calendar(record: Dict[str, Any], output_path: str,
                            final_title: str = "", video_path: str = "",
                            target_duration: str = "1h00 - 1h30",
                            status: str = "Đã xuất video",
                            created_at: Optional[datetime] = None) -> str:
    """Thêm/cập nhật một dòng lịch nội dung; cùng ID/video không bị nhân đôi."""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để xuất lịch nội dung.") from exc

    target = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    if os.path.isfile(target) and os.path.getsize(target) > 0 and zipfile.is_zipfile(target):
        wb = load_workbook(target)
    else:
        # Một lần save bị ngắt hoặc Excel/antivirus khóa file có thể để lại file
        # 0 byte/không còn là ZIP. Giữ bản hỏng để chẩn đoán rồi tự tạo lịch mới.
        if os.path.isfile(target):
            stem, ext = os.path.splitext(target)
            backup = f"{stem}.corrupt_{time.strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
            os.replace(target, backup)
        wb = Workbook()
    if "Lịch nội dung" in wb.sheetnames:
        sheet = wb["Lịch nội dung"]
    else:
        sheet = wb.active
        sheet.title = "Lịch nội dung"
        sheet.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])
    if sheet.max_row < 1 or sheet.cell(1, 1).value != "STT":
        sheet.delete_rows(1, max(1, sheet.max_row))
        sheet.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])

    record_id = str(record.get("id") or "").strip()
    absolute_video = os.path.abspath(video_path) if video_path else ""
    row_index = 0
    for index in range(2, sheet.max_row + 1):
        if ((record_id and str(sheet.cell(index, 10).value or "") == record_id) or
                (absolute_video and str(sheet.cell(index, 11).value or "") == absolute_video)):
            row_index = index
            break
    if not row_index:
        row_index = max(2, sheet.max_row + 1)
    moment = created_at or datetime.now()
    values = [
        f"=ROW()-1", moment.date(),
        (f'=CHOOSE(WEEKDAY(B{row_index},2),"Thứ Hai","Thứ Ba","Thứ Tư",'
         '"Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ Nhật")'),
        str(record.get("best_publish_time") or "20:00"), _calendar_topic(record),
        _calendar_title(record, final_title), str(target_duration or "1h00 - 1h30"),
        str(record.get("source") or "Nhập trực tiếp"),
        _calendar_status({"status": status}, default="Đã xuất video"),
        record_id, absolute_video,
    ]
    for col, value in enumerate(values, 1):
        sheet.cell(row_index, col, value)
    _style_content_calendar(sheet, sheet.max_row)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    # openpyxl ghi trực tiếp sẽ truncate file đích trước khi hoàn tất. Ghi sang
    # file tạm cùng thư mục rồi replace để một lần dừng app không phá workbook.
    handle, temporary = tempfile.mkstemp(
        prefix=".content-calendar-", suffix=".xlsx", dir=os.path.dirname(target))
    os.close(handle)
    try:
        wb.save(temporary)
        try:
            os.replace(temporary, target)
        except PermissionError:
            stem, ext = os.path.splitext(target)
            target = f"{stem}_{time.strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
            os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return target


def export_excel(records: Sequence[Dict[str, Any]], output_path: str) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu openpyxl. Hãy chạy CAI_DAT.bat hoặc pip install openpyxl.") from exc

    rows = [dict(record) for record in records]
    if not rows:
        raise ValueError("Không có ý tưởng nào để xuất Excel.")
    target = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    wb = Workbook()
    plan = wb.active
    plan.title = "Kế hoạch sản xuất"
    dashboard = wb.create_sheet("Dashboard", 0)
    calendar = wb.create_sheet("Lịch nội dung", 1)
    series_sheet = wb.create_sheet("Cụm Series")
    raw = wb.create_sheet("Dữ liệu gốc")
    config = wb.create_sheet("Cấu hình")
    prompts = wb.create_sheet("Prompt mẫu")

    navy, cyan, teal = "0B1626", "14D8D4", "0F766E"
    white, muted, pale = "FFFFFF", "60758A", "E8F7F6"
    green, amber, red = "DCFCE7", "FEF3C7", "FEE2E2"
    thin = Side(style="thin", color="CBD5E1")
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(color=white, bold=True)

    calendar.append(CONTENT_CALENDAR_HEADERS + ["ID nội dung", "Đường dẫn video"])
    for row_index, record in enumerate(rows, 2):
        titles = list(record.get("titles") or [])
        calendar.append([
            f"=ROW()-1", datetime.now().date(),
            (f'=CHOOSE(WEEKDAY(B{row_index},2),"Thứ Hai","Thứ Ba","Thứ Tư",'
             '"Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ Nhật")'),
            record.get("best_publish_time") or "20:00", _calendar_topic(record),
            _calendar_title(record), "1h00 - 1h30", record.get("source", ""),
            _calendar_status(record), record.get("id", ""), "",
        ])
    _style_content_calendar(calendar, calendar.max_row)

    config_rows = [
        ["Tham số", "Giá trị", "Ý nghĩa"],
        ["Trọng số Hook", .45, "Tỷ trọng trong điểm ưu tiên"],
        ["Trọng số Plot Twist", .40, "Tỷ trọng trong điểm ưu tiên"],
        ["Trọng số Dễ sản xuất", .15, "Cao=10, Trung=6, Thấp=3"],
        ["Ngưỡng Nên làm", 7.2, "Điểm ưu tiên tối thiểu"],
        ["Ngưỡng Chờ đợi", 5.5, "Dưới mức này là Không nên làm"],
        ["Tốc độ đọc (từ/phút)", 150, "Dùng ước lượng thời lượng"],
    ]
    for row in config_rows:
        config.append(row)
    config.freeze_panes = "A2"
    config.column_dimensions["A"].width = 28
    config.column_dimensions["B"].width = 14
    config.column_dimensions["C"].width = 42
    for cell in config[1]:
        cell.fill, cell.font = header_fill, header_font
    for cell in config[2][1:2] + config[3][1:2] + config[4][1:2]:
        cell.number_format = "0%"
        cell.fill = PatternFill("solid", fgColor="FFF7CC")
    for row in config.iter_rows(min_row=2, max_row=7, min_col=2, max_col=2):
        row[0].fill = PatternFill("solid", fgColor="FFF7CC")

    columns = [
        "ID", "Chọn sản xuất", "Tiêu đề gốc", "Nguồn", "URL nguồn", "Tác giả",
        "Ngày đăng", "Thể loại (chính)", "Cảm xúc chủ đạo", "Điểm Hook",
        "Điểm Plot Twist", "Điểm ưu tiên", "Chủ đề", "Series",
        "Từ khóa gợi ý", "Tag gợi ý", "Số từ", "Ước lượng thời lượng (phút)",
        "Dễ dàng sản xuất", "Khuyến nghị", "Thời điểm đăng tối ưu",
    ] + [f"Tiêu đề {i}" for i in range(1, 9)] \
      + [f"Mô tả {i}" for i in range(1, 4)] \
      + [f"Thumbnail {i}" for i in range(1, 4)] \
      + ["Outline (dàn ý 3 phần)", "Các nhân vật chính", "Trạng thái sản xuất",
         "Ghi chú", "Provider AI", "Lỗi AI"]
    last_col = get_column_letter(len(columns))
    plan.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    plan["A1"] = "GỐC MÍT KỂ CHUYỆN · KẾ HOẠCH SẢN XUẤT YOUTUBE"
    plan["A1"].fill = PatternFill("solid", fgColor=navy)
    plan["A1"].font = Font(color=white, bold=True, size=16)
    plan["A1"].alignment = Alignment(horizontal="left", vertical="center")
    plan.row_dimensions[1].height = 32
    plan.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    plan["A2"] = ("Ô vàng là trường có thể chỉnh. Điểm ưu tiên và Khuyến nghị là công thức; "
                   "lọc 'Nên làm' để đưa thẳng sang AutoDubVN.")
    plan["A2"].font = Font(color=muted, italic=True)
    plan["A2"].alignment = Alignment(wrap_text=True)
    for col, label in enumerate(columns, 1):
        cell = plan.cell(4, col, label)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=cyan))
    plan.row_dimensions[4].height = 42

    hook_col, twist_col, priority_col = 10, 11, 12
    ease_col, reco_col = 19, 20
    for idx, record in enumerate(rows, 5):
        titles = list(record.get("titles") or [])
        descriptions = list(record.get("descriptions") or [])
        thumbnails = list(record.get("thumbnails") or [])
        values = [
            record.get("id", ""), "Có" if record.get("selected") else "Không",
            record.get("title_original", ""), record.get("source", ""),
            record.get("source_url", ""), record.get("author", ""),
            record.get("published_at", ""), record.get("primary_genre", ""),
            record.get("emotion", ""), int(record.get("hook_score") or 0),
            int(record.get("plot_twist_score") or 0), None,
            ", ".join(record.get("themes") or []), record.get("series", ""),
            ", ".join(record.get("keywords") or []), ", ".join(record.get("tags") or []),
            int(record.get("word_count") or 0), None,
            record.get("production_ease", "Trung"), None,
            record.get("best_publish_time", ""),
        ]
        values += (titles + [""] * 8)[:8]
        values += (descriptions + [""] * 3)[:3]
        values += ([_thumb_text(x) for x in thumbnails] + [""] * 3)[:3]
        values += [
            record.get("outline", ""), ", ".join(record.get("main_characters") or []),
            record.get("status", "Chưa làm"), record.get("notes", ""),
            record.get("analysis_provider", ""), record.get("analysis_error", ""),
        ]
        for col, value in enumerate(values, 1):
            plan.cell(idx, col, value)
        plan.cell(idx, priority_col,
                  f'=ROUND(J{idx}*\'Cấu hình\'!$B$2+K{idx}*\'Cấu hình\'!$B$3+'
                  f'IF(S{idx}="Cao",10,IF(S{idx}="Trung",6,3))*\'Cấu hình\'!$B$4,1)')
        plan.cell(idx, 18, f'=ROUND(Q{idx}/\'Cấu hình\'!$B$7,1)')
        plan.cell(idx, reco_col,
                  f'=IF(L{idx}>=\'Cấu hình\'!$B$5,"Nên làm",'
                  f'IF(L{idx}>=\'Cấu hình\'!$B$6,"Chờ đợi","Không nên làm"))')
        for col in (2, 8, 9, 10, 11, 13, 14, 15, 16, 19, 21, 38, 39):
            plan.cell(idx, col).fill = PatternFill("solid", fgColor="FFF7CC")
        for col in range(1, len(columns) + 1):
            plan.cell(idx, col).alignment = Alignment(
                vertical="top", wrap_text=col >= 3)
        plan.row_dimensions[idx].height = 48

    plan.freeze_panes = "C5"
    plan.auto_filter.ref = f"A4:{last_col}{len(rows) + 4}"
    table = Table(displayName="ProductionPlan", ref=f"A4:{last_col}{len(rows) + 4}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    plan.add_table(table)
    widths = {
        1: 20, 2: 14, 3: 34, 4: 16, 5: 32, 6: 18, 7: 16, 8: 23, 9: 18,
        10: 11, 11: 13, 12: 12, 13: 28, 14: 30, 15: 30, 16: 30, 17: 11,
        18: 17, 19: 17, 20: 15, 21: 27,
    }
    for col in range(22, 30):
        widths[col] = 42
    for col in range(30, 33):
        widths[col] = 52
    for col in range(33, 36):
        widths[col] = 50
    widths.update({36: 58, 37: 32, 38: 20, 39: 30, 40: 24, 41: 42})
    for col, width in widths.items():
        plan.column_dimensions[get_column_letter(col)].width = width
    max_row = len(rows) + 4
    dv_yes = DataValidation(type="list", formula1='"Có,Không"', allow_blank=False)
    dv_ease = DataValidation(type="list", formula1='"Cao,Trung,Thấp"')
    dv_status = DataValidation(
        type="list", formula1='"Chưa làm,Đang viết,Đã duyệt,Đang dựng,Đã đăng,Tạm hoãn"')
    plan.add_data_validation(dv_yes); dv_yes.add(f"B5:B{max_row}")
    plan.add_data_validation(dv_ease); dv_ease.add(f"S5:S{max_row}")
    plan.add_data_validation(dv_status); dv_status.add(f"AL5:AL{max_row}")
    for rng in (f"J5:J{max_row}", f"K5:K{max_row}", f"L5:L{max_row}"):
        plan.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=1, start_color="FEE2E2",
            mid_type="num", mid_value=6, mid_color="FEF3C7",
            end_type="num", end_value=10, end_color="DCFCE7"))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Nên làm\""],
                                       fill=PatternFill("solid", fgColor=green)))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Chờ đợi\""],
                                       fill=PatternFill("solid", fgColor=amber)))
    plan.conditional_formatting.add(
        f"T5:T{max_row}", FormulaRule(formula=["$T5=\"Không nên làm\""],
                                       fill=PatternFill("solid", fgColor=red)))
    plan["L4"].comment = Comment(
        "Hook × trọng số + Plot Twist × trọng số + mức dễ sản xuất × trọng số.",
        "AutoDubVN")

    # Dữ liệu gốc giữ nguyên để kiểm toán nguồn.
    raw_headers = ["ID", "Tiêu đề gốc", "Nguồn", "URL nguồn", "Ngôn ngữ",
                   "Số từ", "Nội dung đang dùng", "Nội dung thô"]
    raw.append(raw_headers)
    for record in rows:
        raw.append([
            record.get("id", ""), record.get("title_original", ""),
            record.get("source", ""), record.get("source_url", ""),
            record.get("language", ""), int(record.get("word_count") or 0),
            str(record.get("content") or "")[:32700],
            str(record.get("raw_content") or "")[:32700],
        ])
    for cell in raw[1]:
        cell.fill, cell.font = header_fill, header_font
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:H{len(rows) + 1}"
    for col, width in enumerate((20, 38, 16, 35, 12, 12, 85, 85), 1):
        raw.column_dimensions[get_column_letter(col)].width = width
    for row in raw.iter_rows(min_row=2):
        row[6].alignment = row[7].alignment = Alignment(wrap_text=True, vertical="top")
        raw.row_dimensions[row[0].row].height = 72

    series_names = sorted({str(r.get("series") or "Chưa phân cụm") for r in rows})
    series_sheet.append(["Series", "Số ý tưởng", "Hook trung bình", "Twist trung bình",
                         "Chủ đề đại diện"])
    for i, name in enumerate(series_names, 2):
        series_sheet.cell(i, 1, name)
        series_sheet.cell(i, 2, f'=COUNTIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i})')
        series_sheet.cell(i, 3, f'=IFERROR(AVERAGEIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i},\'Kế hoạch sản xuất\'!$J$5:$J${max_row}),0)')
        series_sheet.cell(i, 4, f'=IFERROR(AVERAGEIF(\'Kế hoạch sản xuất\'!$N$5:$N${max_row},A{i},\'Kế hoạch sản xuất\'!$K$5:$K${max_row}),0)')
        example = next((r for r in rows if str(r.get("series") or "Chưa phân cụm") == name), {})
        series_sheet.cell(i, 5, ", ".join(example.get("themes") or []))
    for cell in series_sheet[1]:
        cell.fill, cell.font = header_fill, header_font
    series_sheet.freeze_panes = "A2"
    for col, width in enumerate((38, 15, 18, 18, 42), 1):
        series_sheet.column_dimensions[get_column_letter(col)].width = width

    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:H2")
    dashboard["A1"] = "BẢNG ƯU TIÊN SẢN XUẤT NỘI DUNG"
    dashboard["A1"].fill = PatternFill("solid", fgColor=navy)
    dashboard["A1"].font = Font(color=white, bold=True, size=18)
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")
    cards = [
        ("A4", "Tổng ý tưởng", f'=COUNTA(\'Kế hoạch sản xuất\'!$A$5:$A${max_row})'),
        ("C4", "Nên làm", f'=COUNTIF(\'Kế hoạch sản xuất\'!$T$5:$T${max_row},"Nên làm")'),
        ("E4", "Hook TB", f'=ROUND(AVERAGE(\'Kế hoạch sản xuất\'!$J$5:$J${max_row}),1)'),
        ("G4", "Twist TB", f'=ROUND(AVERAGE(\'Kế hoạch sản xuất\'!$K$5:$K${max_row}),1)'),
    ]
    for anchor, label, formula in cards:
        col = dashboard[anchor].column
        dashboard.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        dashboard.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 1)
        dashboard.cell(4, col, label)
        dashboard.cell(5, col, formula)
        dashboard.cell(4, col).fill = PatternFill("solid", fgColor=teal)
        dashboard.cell(4, col).font = Font(color=white, bold=True)
        dashboard.cell(5, col).fill = PatternFill("solid", fgColor=pale)
        dashboard.cell(5, col).font = Font(color=navy, bold=True, size=20)
        dashboard.cell(5, col).alignment = Alignment(horizontal="center", vertical="center")
    dashboard["A8"] = "Cách dùng"
    dashboard["A8"].font = Font(bold=True, size=13, color=navy)
    dashboard.merge_cells("A9:H11")
    dashboard["A9"] = (
        "1) Lọc Khuyến nghị = Nên làm. 2) Chọn một Tiêu đề AI và chỉnh ô vàng. "
        "3) Đổi Chọn sản xuất = Có. 4) Trong AutoDubVN bấm Dùng trong AI Story "
        "để chuyển tiêu đề, outline, tag và mô tả sang quy trình dựng video.")
    dashboard["A9"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard["A9"].fill = PatternFill("solid", fgColor="F8FAFC")
    dashboard["A9"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, 9):
        dashboard.column_dimensions[get_column_letter(col)].width = 16
    if series_names:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Số ý tưởng theo Series"
        chart.y_axis.title = "Series"
        chart.x_axis.title = "Số ý tưởng"
        data = Reference(series_sheet, min_col=2, min_row=1,
                         max_row=len(series_names) + 1)
        cats = Reference(series_sheet, min_col=1, min_row=2,
                         max_row=len(series_names) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 8, 17
        dashboard.add_chart(chart, "A13")

    prompts.sheet_view.showGridLines = False
    prompts.append(["Nhiệm vụ", "Prompt mẫu"])
    prompts.append(["Đánh giá + kế hoạch đa tầng", ANALYSIS_PROMPT])
    prompts.append(["Tạo 8 tiêu đề", TITLE_PROMPT])
    prompts.append(["Tạo 3 mô tả", DESCRIPTION_PROMPT])
    for cell in prompts[1]:
        cell.fill, cell.font = header_fill, header_font
    prompts.column_dimensions["A"].width = 32
    prompts.column_dimensions["B"].width = 110
    for row in prompts.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        prompts.row_dimensions[row[0].row].height = 150

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(target)
    return target


def export_plan(records: Sequence[Dict[str, Any]], output_dir: str = "",
                stem: str = "ke_hoach_noi_dung") -> Dict[str, str]:
    folder = os.path.abspath(output_dir or DEFAULT_OUTPUT_DIR)
    os.makedirs(folder, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^\w\-]+", "_", stem, flags=re.UNICODE).strip("_")
    xlsx = export_excel(records, os.path.join(folder, f"{safe_stem}_{stamp}.xlsx"))
    json_path = os.path.join(folder, f"{safe_stem}_{stamp}.json")
    with open(json_path, "w", encoding="utf-8") as handle:
        json.dump(list(records), handle, ensure_ascii=False, indent=2)
    return {"xlsx": xlsx, "json": json_path, "output_dir": folder}


def sample_records() -> List[Dict[str, Any]]:
    samples = [
        {"titleOriginal": "Ba người con tranh căn nhà của mẹ già",
         "source": "VnExpress Tâm sự", "sourceUrl": "https://vnexpress.net/tam-su",
         "language": "vi", "rawContent": (
             "Bà Hòa đã ngoài bảy mươi, sống một mình trong căn nhà cấp bốn. "
             "Ba người con chỉ trở về khi nghe tin bà định lập di chúc. Trong bữa cơm giỗ, "
             "một tờ giấy cũ bất ngờ xuất hiện khiến mọi người tranh cãi dữ dội. ") * 120},
        {"titleOriginal": "婆媳矛盾：老人被赶出家门",
         "source": "Zhihu 盐选", "sourceUrl": "https://www.zhihu.com/",
         "language": "zh", "rawContent": (
             "婆婆和儿媳因为赡养和房产发生矛盾，老人被赶出家门。后来一份遗产文件揭开真相。") * 280},
        {"titleOriginal": "AITA for refusing to pay my father's nursing home?",
         "source": "Reddit AITA", "sourceUrl": "https://www.reddit.com/r/AmItheAsshole",
         "language": "en", "rawContent": (
             "My siblings argued about who should care for our elderly father after our mother died. "
             "Everyone wanted the inheritance, but nobody agreed to visit him at the nursing home. ") * 170},
    ]
    return [analyze_record(normalize_record(row, i), use_ai=False)
            for i, row in enumerate(samples)]
